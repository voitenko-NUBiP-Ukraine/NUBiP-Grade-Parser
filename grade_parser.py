#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Grade Report Parser: converts a saved Moodle grader report (.mhtml / .html) to TXT and Excel.

Run without arguments for the GUI, or drop files onto the script / built .exe.
Results go to <chosen folder>/<Title>_<YYYY-MM-DD>_<HH-MM-SS>/
"""

from __future__ import annotations

import email
import email.policy
import math
import re
import shutil
import sys
import textwrap
import traceback
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    HAVE_RICH_TEXT = True
except Exception:
    HAVE_RICH_TEXT = False

APP_NAME = "Grade Report Parser"
SUPPORTED_SUFFIXES = {".mhtml", ".mht", ".html", ".htm"}

ENABLE_HEATMAP = True


@dataclass
class GradeItem:

    item_id: str
    name: str
    group: str
    kind: str
    is_total: bool
    max_points: Optional[float]
    average: str
    url: str


@dataclass
class Student:
    name: str
    fields: List[str]
    grades: List[str]


@dataclass
class GradeReport:
    title: str
    source_name: str
    field_headers: List[str]
    items: List[GradeItem]
    students: List[Student]
    average_label: str


def _decode(data: bytes, hinted: Optional[str] = None) -> str:
    candidates: List[str] = []
    if hinted:
        candidates.append(hinted)
    match = re.search(rb"charset=[\"']?([\w-]+)", data[:8192], re.I)
    if match:
        candidates.append(match.group(1).decode("ascii", "ignore"))
    candidates += ["utf-8", "cp1251"]
    for enc in candidates:
        try:
            return data.decode(enc)
        except (LookupError, UnicodeDecodeError):
            continue
    return data.decode("utf-8", errors="replace")


def load_html(path: Path) -> str:
    raw = path.read_bytes()
    is_mhtml = path.suffix.lower() in {".mhtml", ".mht"} or b"MIME-Version" in raw[:4096]

    if not is_mhtml:
        return _decode(raw)

    message = email.message_from_bytes(raw, policy=email.policy.compat32)
    candidates: List[str] = []
    for part in message.walk():
        if part.get_content_type() != "text/html":
            continue
        payload = part.get_payload(decode=True)
        if payload:
            candidates.append(_decode(payload, part.get_content_charset()))

    for html in candidates:
        if "gradeparent" in html:
            return html
    if candidates:
        return candidates[0]
    raise ValueError("No HTML part found inside the MHTML file.")


_NUMBER_RE = re.compile(r"^[-+]?\d+(?:[.,]\d+)?$")
_RANGE_MAX_RE = re.compile(r"[-\u2013\u2014]\s*(\d+(?:[.,]\d+)?)\s*$")
_ITEM_CLASS_RE = re.compile(r"^i(\d+)$")


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _text(el) -> str:
    return _clean(el.get_text(" ", strip=True)) if el is not None else ""


def _first_string(el) -> str:
    for s in el.stripped_strings:
        return _clean(s)
    return ""


def to_number(text: str):
    if _NUMBER_RE.match(text):
        value = float(text.replace(",", "."))
        return int(value) if value.is_integer() else value
    return None


def fmt_num(value: Optional[float]) -> str:
    if value is None:
        return ""
    return str(int(value)) if float(value).is_integer() else f"{value:g}"


def _span(cell) -> int:
    try:
        return max(1, int(cell.get("colspan", 1)))
    except (TypeError, ValueError):
        return 1


def _class_set(el) -> set:
    return set(el.get("class") or [])


def _by_item_id(row) -> dict:
    result = {}
    for td in row.find_all("td"):
        for cls in td.get("class") or []:
            m = _ITEM_CLASS_RE.match(cls)
            if m:
                result[m.group(1)] = td
                break
    return result


def _category_labels(row, n_items: int, n_user_cols: int) -> List[str]:
    labels: List[str] = []
    for cell in row.find_all(["th", "td"], recursive=False):
        label = ""
        if cell.name == "th":
            header = cell.select_one(".gradeitemheader")
            label = _text(header) if header else ""
        labels.extend([label] * _span(cell))

    if len(labels) == n_items + n_user_cols:
        labels = labels[n_user_cols:]
    labels = (labels + [""] * n_items)[:n_items]
    return labels


def parse_report(html: str, source_name: str) -> GradeReport:
    soup = BeautifulSoup(html, "html.parser")

    h1 = soup.select_one("h1.h2.mb-0") or soup.find("h1")
    title = _text(h1)
    if not title:
        title = _text(soup.title) or Path(source_name).stem

    parent = soup.find("div", class_="gradeparent")
    if parent is None:
        raise ValueError('Grade table not found (<div class="gradeparent"> is missing).')
    table = parent.find("table")
    if table is None:
        raise ValueError("Grade table not found inside the gradeparent block.")

    rows = table.find_all("tr")
    heading = next((r for r in rows if "heading" in _class_set(r)), None)
    if heading is None:
        raise ValueError("Column header row not found in the grade table.")
    heading_idx = rows.index(heading)
    category_rows = rows[:heading_idx]

    field_keys: List[str] = []
    field_headers: List[str] = []
    item_cells = []
    for th in heading.find_all("th", recursive=False):
        classes = _class_set(th)
        if "userfield" in classes:
            key = next(
                (c for c in th.get("class") if c not in {"userfield", "cell"} and not re.match(r"^c\d+$", c)),
                "",
            )
            field_keys.append(key)
            field_headers.append(_first_string(th) or key)
        elif classes & {"item", "courseitem", "categoryitem"}:
            item_cells.append(th)

    n_items = len(item_cells)
    if n_items == 0:
        raise ValueError("No grade columns found in the table.")
    n_user_cols = len(field_keys) + 1

    group_rows = [_category_labels(r, n_items, n_user_cols) for r in category_rows[1:]]

    range_cells, avg_cells, average_label = {}, {}, "Average"
    for row in rows[heading_idx + 1:]:
        classes = _class_set(row)
        if "range" in classes and not range_cells:
            range_cells = _by_item_id(row)
        elif "avg" in classes and not avg_cells:
            avg_cells = _by_item_id(row)
            th = row.find("th")
            if th is not None:
                average_label = _first_string(th) or average_label

    items: List[GradeItem] = []
    for idx, th in enumerate(item_cells):
        item_id = th.get("data-itemid", "")
        header = th.select_one(".gradeitemheader")
        name = _text(header) or f"Item {idx + 1}"

        icon = th.select_one("img.itemicon")
        if icon is not None and icon.get("alt"):
            kind = _clean(icon["alt"])
        else:
            icon = th.select_one("i.itemicon")
            kind = _clean(icon.get("aria-label", "")) if icon is not None else ""

        path = []
        for labels in group_rows:
            if labels[idx] and labels[idx] not in path:
                path.append(labels[idx])

        max_points = None
        if item_id in range_cells:
            m = _RANGE_MAX_RE.search(_text(range_cells[item_id]))
            if m:
                max_points = float(m.group(1).replace(",", "."))

        items.append(
            GradeItem(
                item_id=item_id,
                name=name,
                group=" \u203a ".join(path),
                kind=kind,
                is_total="courseitem" in _class_set(th) or "categoryitem" in _class_set(th),
                max_points=max_points,
                average=_text(avg_cells[item_id]) if item_id in avg_cells else "",
                url=(header.get("href", "") if header is not None and header.name == "a" else ""),
            )
        )

    students: List[Student] = []
    for row in rows:
        if "userrow" not in _class_set(row):
            continue
        name_th = row.find("th")
        link = name_th.select_one("a.username") if name_th is not None else None
        name = _text(link) if link is not None else (_first_string(name_th) if name_th is not None else "")

        values = []
        for key in field_keys:
            cell = row.find("td", class_=key) if key else None
            values.append(_text(cell))

        by_id = {td.get("data-itemid"): td for td in row.find_all("td") if td.get("data-itemid")}
        grades = []
        for item in items:
            td = by_id.get(item.item_id)
            if td is None:
                grades.append("-")
                continue
            span = td.select_one(".gradevalue")
            value = _text(span) if span is not None else _first_string(td)
            grades.append(value or "-")

        students.append(Student(name=name, fields=values, grades=grades))

    return GradeReport(
        title=title,
        source_name=source_name,
        field_headers=field_headers,
        items=items,
        students=students,
        average_label=average_label,
    )


def export_txt(report: GradeReport, path: Path, parsed_at: datetime) -> None:
    items, students = report.items, report.students

    def cut(text: str, w: int) -> str:
        return text if len(text) <= w else text[: w - 1] + "\u2026"

    columns = []
    columns.append((["#"], max(2, len(str(len(students)))), "right",
                    [str(i) for i in range(1, len(students) + 1)], "", ""))

    name_w = min(40, max([len("Student")] + [len(s.name) for s in students]))
    columns.append((["Student"], name_w, "left", [cut(s.name, name_w) for s in students],
                    "Max points", report.average_label))

    for j, item in enumerate(items):
        longest_word = max((len(w) for w in item.name.split()), default=1)
        wrap_w = max(longest_word, min(22, max(8, math.ceil(len(item.name) / 6) + 2)))
        header = textwrap.wrap(item.name, wrap_w, break_long_words=False) or [item.name]
        values = [s.grades[j] for s in students]
        max_text, avg_text = fmt_num(item.max_points), item.average or ""
        width = max([len(line) for line in header] + [len(v) for v in values] + [len(max_text), len(avg_text)])
        columns.append((header, width, "center", values, max_text, avg_text))

    for k, header_text in enumerate(report.field_headers):
        values = [s.fields[k] for s in students]
        width = max([len(header_text)] + [len(v) for v in values])
        columns.append(([header_text], width, "left", values, "", ""))

    def align(text: str, w: int, mode: str) -> str:
        return text.rjust(w) if mode == "right" else text.center(w) if mode == "center" else text.ljust(w)

    def make_line(cells: List[str]) -> str:
        return " \u2502 ".join(align(c, col[1], col[2]) for c, col in zip(cells, columns)).rstrip()

    sep = "\u2500\u253c\u2500".join("\u2500" * col[1] for col in columns)

    out: List[str] = []
    table_width = len(sep)
    bar = "=" * min(max(table_width, len(report.title) + 4), 120)
    out += [bar, f"  {report.title}", bar, ""]

    n_lines = max(len(col[0]) for col in columns)
    for line_no in range(n_lines):
        cells = []
        for col in columns:
            pad = n_lines - len(col[0])
            idx = line_no - pad
            cells.append(col[0][idx] if idx >= 0 else "")
        out.append(make_line(cells))
    out.append(sep)
    out.append(make_line([col[4] for col in columns]))
    out.append(sep)
    for i in range(len(students)):
        out.append(make_line([col[3][i] for col in columns]))
    out.append(sep)
    out.append(make_line([col[5] for col in columns]))
    out.append("")

    path.write_text("\n".join(out), encoding="utf-8-sig")


C_DARK = "1B4332"
C_GREEN_A = "2D6A4F"
C_GREEN_B = "40916C"
C_HEADER_BG = "E8F3EC"
C_ZEBRA = "F6FAF7"
C_GOLD_BG = "FFF4D6"
C_GOLD_HEADER = "F6E3AE"
C_GRID = "DDE5E0"
C_MUTED = "9AA5A0"
FONT = "Calibri"


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _style_range(ws, row: int, col_from: int, col_to: int, fill: PatternFill = None, border: Border = None):
    for col in range(col_from, col_to + 1):
        cell = ws.cell(row=row, column=col)
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border


def _header_value(name: str, max_text: str):
    if not HAVE_RICH_TEXT:
        return name + (f"\nMax {max_text}" if max_text else "")
    blocks = [TextBlock(InlineFont(rFont=FONT, sz=10, b=True, color=C_DARK), name)]
    if max_text:
        blocks.append(TextBlock(InlineFont(rFont=FONT, sz=9, i=True, color="6B7F75"), f"\nMax {max_text}"))
    return CellRichText(*blocks)


def export_xlsx(report: GradeReport, path: Path, parsed_at: datetime) -> None:
    items, students = report.items, report.students
    n, nf = len(items), len(report.field_headers)

    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    ws.sheet_properties.tabColor = C_GREEN_A
    ws.sheet_view.showGridLines = False

    c_num, c_name, c_item0 = 1, 2, 3
    c_itemN = c_item0 + n - 1
    c_det0 = c_itemN + 1
    c_last = c_itemN + nf

    r_title, r_sub, r_gap, r_group, r_head = 1, 2, 3, 4, 5
    r_first = r_head + 1
    r_last = r_first + len(students) - 1
    r_avg = r_last + 1

    item_w = 17
    total_w = 15

    ws.merge_cells(start_row=r_title, start_column=1, end_row=r_title, end_column=c_last)
    _style_range(ws, r_title, 1, c_last, fill=_fill(C_DARK))
    cell = ws.cell(row=r_title, column=1, value=report.title)
    cell.font = Font(name=FONT, size=20, bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r_title].height = 40

    ws.merge_cells(start_row=r_sub, start_column=1, end_row=r_sub, end_column=c_last)
    cell = ws.cell(
        row=r_sub,
        column=1,
        value=(
            f"Grader report   \u2022   Students: {len(students)}   \u2022   Columns: {n}   "
            f"\u2022   Parsed: {parsed_at:%Y-%m-%d %H:%M}   \u2022   Source: {report.source_name}"
        ),
    )
    cell.font = Font(name=FONT, size=10, italic=True, color="5B6B63")
    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[r_sub].height = 22
    ws.row_dimensions[r_gap].height = 8

    white_bold = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def put_group(col_from: int, col_to: int, label: str, color: str):
        if col_to > col_from:
            ws.merge_cells(start_row=r_group, start_column=col_from, end_row=r_group, end_column=col_to)
        _style_range(ws, r_group, col_from, col_to, fill=_fill(color))
        c = ws.cell(row=r_group, column=col_from, value=label)
        c.font = white_bold
        c.alignment = center_wrap

    put_group(c_num, c_name, "Student", C_DARK)

    seg_start, seg_label, palette_i = 0, "", 0
    palette = [C_GREEN_A, C_GREEN_B]
    segments: List[Tuple[int, int, str]] = []
    for i, item in enumerate(items + [None]):
        label = item.group if item is not None else ""
        if i == 0:
            seg_label = label
        elif label != seg_label or item is None:
            segments.append((seg_start, i - 1, seg_label))
            seg_start, seg_label = i, label
    for a, b, label in segments:
        if not label:
            continue
        put_group(c_item0 + a, c_item0 + b, label, palette[palette_i % 2])
        palette_i += 1
    if nf:
        put_group(c_det0, c_last, "Student details", C_DARK)
    ws.row_dimensions[r_group].height = 30

    head_border = Border(bottom=Side(style="medium", color=C_GREEN_A))
    _style_range(ws, r_head, 1, c_last, fill=_fill(C_HEADER_BG), border=head_border)

    def head_cell(col: int, value, left: bool = False, fill: str = None):
        c = ws.cell(row=r_head, column=col, value=value)
        c.font = Font(name=FONT, size=10, bold=True, color=C_DARK)
        c.alignment = Alignment(
            horizontal="left" if left else "center", vertical="center", wrap_text=True, indent=1 if left else 0
        )
        c.border = head_border
        if fill:
            c.fill = _fill(fill)

    head_cell(c_num, "#")
    head_cell(c_name, "Student", left=True)

    max_header_height = 48.0
    for j, item in enumerate(items):
        col = c_item0 + j
        w = total_w if item.is_total else item_w
        ws.column_dimensions[get_column_letter(col)].width = w
        head_cell(col, _header_value(item.name, fmt_num(item.max_points)), fill=C_GOLD_HEADER if item.is_total else None)
        lines = math.ceil(len(item.name) * 1.15 / w) + (1 if item.max_points is not None else 0)
        max_header_height = max(max_header_height, lines * 13.5 + 10)
    ws.row_dimensions[r_head].height = min(190.0, max_header_height)

    for k, header in enumerate(report.field_headers):
        head_cell(c_det0 + k, header, left=True)

    thin = Side(style="thin", color=C_GRID)
    row_border = Border(bottom=thin)
    muted_font = Font(name=FONT, size=11, color=C_MUTED)
    grade_font = Font(name=FONT, size=11)
    bold_font = Font(name=FONT, size=11, bold=True)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", indent=1)

    for i, s in enumerate(students):
        r = r_first + i
        zebra = _fill(C_ZEBRA) if i % 2 else None
        _style_range(ws, r, 1, c_last, fill=zebra, border=row_border)

        c = ws.cell(row=r, column=c_num, value=i + 1)
        c.font, c.alignment = Font(name=FONT, size=10, color=C_MUTED), center

        c = ws.cell(row=r, column=c_name, value=s.name)
        c.font, c.alignment = Font(name=FONT, size=11, bold=True, color=C_DARK), left

        for j, raw in enumerate(s.grades):
            item = items[j]
            c = ws.cell(row=r, column=c_item0 + j)
            number = to_number(raw)
            if number is not None:
                c.value = number
                c.number_format = "General"
                c.font = bold_font if item.is_total else grade_font
            else:
                c.value = "\u2013" if raw in ("-", "") else raw
                c.font = muted_font if raw in ("-", "") else grade_font
            c.alignment = center
            if item.is_total:
                c.border = Border(bottom=thin, left=Side(style="thin", color=C_GREEN_B))
                if zebra is None:
                    c.fill = _fill(C_GOLD_BG)

        for k, value in enumerate(s.fields):
            c = ws.cell(row=r, column=c_det0 + k, value=value)
            c.font, c.alignment = Font(name=FONT, size=10, color="44514A"), left

        ws.row_dimensions[r].height = 22

    if students:
        top = Border(top=Side(style="medium", color=C_GREEN_A), bottom=Side(style="thin", color=C_GREEN_A))
        _style_range(ws, r_avg, 1, c_last, fill=_fill(C_HEADER_BG), border=top)
        ws.merge_cells(start_row=r_avg, start_column=c_num, end_row=r_avg, end_column=c_name)
        c = ws.cell(row=r_avg, column=c_num, value=report.average_label)
        c.font = Font(name=FONT, size=11, bold=True, italic=True, color=C_DARK)
        c.alignment = left
        for j, item in enumerate(items):
            c = ws.cell(row=r_avg, column=c_item0 + j, value=item.average or "\u2013")
            c.font = Font(name=FONT, size=10, bold=True, color=C_DARK if item.average else C_MUTED)
            c.alignment = center
        ws.row_dimensions[r_avg].height = 24

        ws.merge_cells(start_row=r_avg + 1, start_column=1, end_row=r_avg + 1, end_column=c_itemN)
        c = ws.cell(
            row=r_avg + 1,
            column=1,
            value="Average row: value (number of graded entries).  \u2013 = no grade.",
        )
        c.font = Font(name=FONT, size=9, italic=True, color=C_MUTED)
        c.alignment = left

    if ENABLE_HEATMAP and students:
        for j, item in enumerate(items):
            letter = get_column_letter(c_item0 + j)
            rng = f"{letter}{r_first}:{letter}{r_last}"
            if item.max_points:
                rule = ColorScaleRule(
                    start_type="num", start_value=0, start_color="F8D9D6",
                    mid_type="num", mid_value=item.max_points / 2, mid_color="FFF3CF",
                    end_type="num", end_value=item.max_points, end_color="CBEBD3",
                )
            else:
                rule = ColorScaleRule(
                    start_type="min", start_color="F8D9D6",
                    mid_type="percentile", mid_value=50, mid_color="FFF3CF",
                    end_type="max", end_color="CBEBD3",
                )
            ws.conditional_formatting.add(rng, rule)

    longest_name = max([len(s.name) for s in students] + [len("Student")])
    ws.column_dimensions[get_column_letter(c_num)].width = 5
    ws.column_dimensions[get_column_letter(c_name)].width = min(44, max(26, longest_name + 4))
    for k in range(nf):
        longest = max([len(s.fields[k]) for s in students] + [len(report.field_headers[k])])
        ws.column_dimensions[get_column_letter(c_det0 + k)].width = min(55, max(14, longest + 3))

    ws.freeze_panes = ws.cell(row=r_first, column=c_item0)
    if students:
        ws.auto_filter.ref = f"A{r_head}:{get_column_letter(c_last)}{r_last}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{r_group}:{r_head}"
    ws.print_area = f"A1:{get_column_letter(c_itemN)}{r_avg + 1 if students else r_head}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "&P / &N"

    wb.properties.title = report.title
    wb.properties.creator = APP_NAME
    wb.save(path)


def resource_path(name: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / name


def safe_name(text: str, limit: int = 80) -> str:
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text)
    text = re.sub(r"\s+", "_", text.strip()).strip("._ ")
    return (text[:limit].rstrip("._ ")) or "report"


def unique_dir(path: Path) -> Path:
    candidate, k = path, 2
    while candidate.exists():
        candidate = path.with_name(f"{path.name}_{k}")
        k += 1
    return candidate


def process_file(
    src: Path,
    output_root: Path,
    want_xlsx: bool = True,
    want_txt: bool = True,
    copy_source: bool = False,
    log: Callable[[str], None] = print,
) -> Tuple[Path, GradeReport]:
    src = Path(src)
    log(f"Reading: {src.name}")
    report = parse_report(load_html(src), src.name)
    log(f"Parsed '{report.title}': {len(report.students)} students, {len(report.items)} columns")

    now = datetime.now()
    base = f"{safe_name(report.title)}_{now:%Y-%m-%d_%H-%M-%S}"
    folder = unique_dir(Path(output_root) / base)
    folder.mkdir(parents=True, exist_ok=True)

    if want_txt:
        export_txt(report, folder / f"{folder.name}.txt", now)
        log("  \u2192 TXT saved")
    if want_xlsx:
        export_xlsx(report, folder / f"{folder.name}.xlsx", now)
        log("  \u2192 Excel saved")
    if copy_source:
        shutil.copy2(src, folder / src.name)
        log("  \u2192 Source page copied")
    return folder, report


def run_gui(initial_files: List[str]) -> int:
    from PyQt5.QtCore import Qt, QThread, QTimer, QUrl, pyqtSignal
    from PyQt5.QtGui import QDesktopServices, QFont, QIcon
    from PyQt5.QtWidgets import (
        QApplication, QCheckBox, QFileDialog, QFrame, QHBoxLayout, QLabel, QListWidget,
        QListWidgetItem, QPlainTextEdit, QProgressBar, QPushButton, QVBoxLayout, QWidget,
    )

    STYLE = """
    QWidget { font-family: 'Segoe UI', 'Noto Sans', Arial; font-size: 10pt; color: #1f2a24; }
    QWidget#root { background: #f3f6f4; }
    QLabel#title { font-size: 20pt; font-weight: 700; color: #1b4332; }
    QLabel#subtitle { color: #6b7f75; }
    QLabel#section { font-weight: 600; color: #2d6a4f; }
    QFrame#drop { background: #ffffff; border: 2px dashed #95b8a5; border-radius: 14px; }
    QFrame#drop[active="true"] { background: #e6f4ec; border: 2px dashed #2d6a4f; }
    QLabel#dropTitle { font-size: 13pt; font-weight: 600; color: #1b4332; }
    QLabel#dropHint { color: #6b7f75; }
    QPushButton { background: #ffffff; border: 1px solid #c5d3cb; border-radius: 8px; padding: 7px 16px; }
    QPushButton:hover { background: #eef5f0; border-color: #95b8a5; }
    QPushButton:disabled { color: #a3aea8; background: #f0f3f1; }
    QPushButton#primary { background: #2d6a4f; color: white; border: none; font-weight: 600; padding: 9px 22px; }
    QPushButton#primary:hover { background: #40916c; }
    QPushButton#primary:disabled { background: #a9c7b8; color: white; }
    QListWidget, QPlainTextEdit { background: #ffffff; border: 1px solid #d5e0d9; border-radius: 10px; padding: 4px; }
    QListWidget::item { padding: 7px 8px; border-radius: 6px; }
    QListWidget::item:selected { background: #e1f0e7; color: #1b4332; }
    QPlainTextEdit { font-family: Consolas, 'DejaVu Sans Mono', monospace; font-size: 9pt; color: #44514a; }
    QProgressBar { background: #dfe8e2; border: none; border-radius: 4px; height: 8px; }
    QProgressBar::chunk { background: #40916c; border-radius: 4px; }
    QCheckBox { spacing: 8px; }
    """

    class Worker(QThread):
        message = pyqtSignal(str)
        done_one = pyqtSignal(str, str, str)
        progress = pyqtSignal(int, int)

        def __init__(self, files, root, xlsx, txt, copy_src):
            super().__init__()
            self.files, self.root, self.xlsx, self.txt, self.copy_src = files, root, xlsx, txt, copy_src

        def run(self):
            total = len(self.files)
            self.progress.emit(0, total)
            for i, f in enumerate(self.files, 1):
                try:
                    folder, _ = process_file(
                        Path(f), self.root, self.xlsx, self.txt, self.copy_src, log=self.message.emit
                    )
                    self.done_one.emit("ok", Path(f).name, str(folder))
                except Exception as exc:
                    self.message.emit(f"ERROR in {Path(f).name}: {exc}")
                    self.message.emit(traceback.format_exc(limit=3))
                    self.done_one.emit("error", Path(f).name, str(exc))
                self.progress.emit(i, total)

    class Window(QWidget):
        def __init__(self):
            super().__init__()
            self.setObjectName("root")
            self.setWindowTitle(APP_NAME)
            self.setAcceptDrops(True)
            self.resize(760, 700)
            self.worker: Optional[Worker] = None
            self.jobs: List[Tuple[List[str], Path]] = []
            self.last_dir = str(Path.home())

            root = QVBoxLayout(self)
            root.setContentsMargins(26, 22, 26, 22)
            root.setSpacing(12)

            title = QLabel(APP_NAME)
            title.setObjectName("title")
            subtitle = QLabel("Convert saved Moodle grade report pages into TXT and Excel files")
            subtitle.setObjectName("subtitle")
            root.addWidget(title)
            root.addWidget(subtitle)

            self.drop = QFrame()
            self.drop.setObjectName("drop")
            self.drop.setProperty("active", False)
            dl = QVBoxLayout(self.drop)
            dl.setContentsMargins(20, 22, 20, 22)
            dl.setSpacing(6)
            dt = QLabel("Drop saved pages here")
            dt.setObjectName("dropTitle")
            dt.setAlignment(Qt.AlignCenter)
            dh = QLabel("Supported: .mhtml, .mht, .html, .htm   \u2022   you will be asked where to save the results")
            dh.setObjectName("dropHint")
            dh.setAlignment(Qt.AlignCenter)
            self.btn_select = QPushButton("Select files\u2026")
            self.btn_select.setObjectName("primary")
            self.btn_select.setCursor(Qt.PointingHandCursor)
            self.btn_select.clicked.connect(self.select_files)
            dl.addWidget(dt)
            dl.addWidget(dh)
            row = QHBoxLayout()
            row.addStretch()
            row.addWidget(self.btn_select)
            row.addStretch()
            dl.addLayout(row)
            root.addWidget(self.drop)

            opts = QHBoxLayout()
            opts.setSpacing(20)
            self.cb_xlsx = QCheckBox("Excel (.xlsx)")
            self.cb_txt = QCheckBox("Text (.txt)")
            self.cb_src = QCheckBox("Copy source page to output folder")
            self.cb_xlsx.setChecked(True)
            self.cb_txt.setChecked(True)
            for cb in (self.cb_xlsx, self.cb_txt, self.cb_src):
                opts.addWidget(cb)
            opts.addStretch()
            root.addLayout(opts)

            self.progress = QProgressBar()
            self.progress.setTextVisible(False)
            self.progress.setFixedHeight(8)
            root.addWidget(self.progress)

            lbl = QLabel("Results  (double-click to open the folder)")
            lbl.setObjectName("section")
            root.addWidget(lbl)
            self.results = QListWidget()
            self.results.setAcceptDrops(False)
            self.results.setMinimumHeight(110)
            self.results.itemDoubleClicked.connect(self.open_result)
            root.addWidget(self.results, 1)

            head = QHBoxLayout()
            lbl = QLabel("Log")
            lbl.setObjectName("section")
            btn_clear = QPushButton("Clear log")
            btn_clear.clicked.connect(lambda: self.log.clear())
            head.addWidget(lbl)
            head.addStretch()
            head.addWidget(btn_clear)
            root.addLayout(head)
            self.log = QPlainTextEdit()
            self.log.setReadOnly(True)
            self.log.setAcceptDrops(False)
            self.log.setMinimumHeight(110)
            root.addWidget(self.log, 1)

        def append_log(self, text: str):
            self.log.appendPlainText(f"[{datetime.now():%H:%M:%S}] {text}")

        def open_result(self, item: QListWidgetItem):
            folder = item.data(Qt.UserRole)
            if folder and Path(folder).exists():
                QDesktopServices.openUrl(QUrl.fromLocalFile(folder))

        def select_files(self):
            files, _ = QFileDialog.getOpenFileNames(
                self, "Select saved pages", "", "Saved pages (*.mhtml *.mht *.html *.htm);;All files (*.*)"
            )
            if files:
                self.add_files(files)

        @staticmethod
        def _dropped_files(event) -> List[str]:
            files = []
            for url in event.mimeData().urls():
                if url.isLocalFile():
                    p = Path(url.toLocalFile())
                    if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES:
                        files.append(str(p))
            return files

        def _set_active(self, active: bool):
            self.drop.setProperty("active", active)
            self.drop.style().unpolish(self.drop)
            self.drop.style().polish(self.drop)

        def dragEnterEvent(self, event):
            if event.mimeData().hasUrls() and self._dropped_files(event):
                event.acceptProposedAction()
                self._set_active(True)
            else:
                event.ignore()

        def dragLeaveEvent(self, event):
            self._set_active(False)

        def dropEvent(self, event):
            self._set_active(False)
            files = self._dropped_files(event)
            if files:
                event.acceptProposedAction()
                QTimer.singleShot(0, lambda: self.add_files(files))

        def add_files(self, files: List[str]):
            if not (self.cb_xlsx.isChecked() or self.cb_txt.isChecked()):
                self.append_log("Nothing to do: enable Excel and/or Text output.")
                return
            folder = QFileDialog.getExistingDirectory(self, "Choose folder to save the results", self.last_dir)
            if not folder:
                self.append_log("Cancelled: no output folder selected.")
                return
            self.last_dir = folder
            self.jobs.append((files, Path(folder)))
            if self.worker is None:
                self.start_next()

        def start_next(self):
            if not self.jobs:
                return
            files, root = self.jobs.pop(0)
            self.btn_select.setEnabled(False)
            self.worker = Worker(files, root, self.cb_xlsx.isChecked(), self.cb_txt.isChecked(), self.cb_src.isChecked())
            self.worker.message.connect(self.append_log)
            self.worker.done_one.connect(self.on_done_one)
            self.worker.progress.connect(self.on_progress)
            self.worker.finished.connect(self.on_finished)
            self.worker.start()

        def on_progress(self, done: int, total: int):
            self.progress.setMaximum(max(total, 1))
            self.progress.setValue(done)

        def on_done_one(self, status: str, name: str, info: str):
            if status == "ok":
                item = QListWidgetItem(f"\u2714  {name}\n     {info}")
                item.setData(Qt.UserRole, info)
                item.setToolTip(info)
            else:
                item = QListWidgetItem(f"\u2716  {name}\n     {info}")
            self.results.insertItem(0, item)

        def on_finished(self):
            self.worker = None
            self.btn_select.setEnabled(True)
            self.start_next()

        def closeEvent(self, event):
            if self.worker is not None:
                self.worker.wait(5000)
            super().closeEvent(event)

    if hasattr(Qt, "AA_EnableHighDpiScaling"):
        QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    if sys.platform == "win32":
        try:
            import ctypes
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("GradeReportParser.App")
        except Exception:
            pass
    app = QApplication(sys.argv)
    app.setStyleSheet(STYLE)
    app.setFont(QFont("Segoe UI", 10))
    icon = QIcon()
    for icon_name in ("icon.ico", "icon.png"):
        icon_path = resource_path(icon_name)
        if icon_path.is_file():
            icon.addFile(str(icon_path))
    app.setWindowIcon(icon)
    window = Window()
    window.setWindowIcon(icon)
    window.show()

    valid = [f for f in initial_files if Path(f).suffix.lower() in SUPPORTED_SUFFIXES and Path(f).is_file()]
    for f in initial_files:
        if f not in valid:
            window.append_log(f"Skipped (unsupported or missing): {f}")
    if valid:
        QTimer.singleShot(150, lambda: window.add_files(valid))

    return app.exec_()


def main() -> int:
    return run_gui(sys.argv[1:])


if __name__ == "__main__":
    sys.exit(main())
